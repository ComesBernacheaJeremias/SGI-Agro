"""Importación desde Excel: plantilla, vista previa con errores por fila, todo o nada."""

from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.modules.identity.authorization import OWNER, READ_ONLY
from tests.factories import login_as
from tests.integration.test_commercial import Shop


def xlsx(header: list[str], rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    sheet = wb.active
    assert sheet is not None
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class Importer(Shop):
    def send(self, key: str, content: bytes, *, confirm: bool = False) -> Any:
        url = f"/api/v1/imports/{key}" + ("" if confirm else "/preview")
        files = {"file": ("datos.xlsx", content, "application/octet-stream")}
        response = self.c.post(url, files=files)
        assert response.status_code == 200, response.json()
        return response.json()


@pytest.fixture
def shop(client: TestClient, db: Session) -> Importer:
    login_as(client, db, OWNER)
    return Importer(client)


PRODUCT_HEADER = ["Código", "Nombre *", "Tipo *", "Unidad *", "Categoría", "IVA %", "Equivalencias"]


def test_imports_are_listed_in_load_order(shop: Importer) -> None:
    keys = [i["key"] for i in shop.c.get("/api/v1/imports").json()]
    assert keys == ["products", "parties", "opening_stock", "opening_balances"]


def test_template_has_headers_and_instructions(shop: Importer) -> None:
    response = shop.c.get("/api/v1/imports/products/template")
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    headers = [c.value for c in wb["Datos"][1]]
    assert headers[:4] == ["Código", "Nombre *", "Tipo *", "Unidad *"]
    assert "Instrucciones" in wb.sheetnames


def test_products_preview_reports_errors_and_saves_nothing(shop: Importer) -> None:
    content = xlsx(
        PRODUCT_HEADER,
        [
            ["", "Glifosato", "Insumo", "L", "Insumos", "21", "0,2 bidón"],
            ["", "Semilla", "insumo", "bolsa-x", "", "", ""],  # unidad inexistente
            ["", "Cajas", "Cualquiera", "un", "", "", ""],  # tipo inválido
            ["", "", "", "", "", "", ""],  # vacía: se ignora
        ],
    )
    result = shop.send("products", content)
    assert (result["total"], result["valid"]) == (3, 1)
    assert [e["row"] for e in result["errors"]] == [3, 4]
    assert "bolsa-x" in result["errors"][0]["message"]
    confirmed = shop.send("products", content, confirm=True)
    assert confirmed["committed"] is False
    names = [
        p["name"] for p in shop.c.get("/api/v1/products", params={"q": "Glifosato"}).json()["items"]
    ]
    assert names == []


def test_products_import_all_or_nothing(shop: Importer) -> None:
    content = xlsx(
        PRODUCT_HEADER,
        [
            ["GLI-1", "Glifosato", "Insumo", "L", "Insumos", "", "0,05 bidón"],
            ["gli-1", "Otro", "Reventa", "un", "", "10,5", ""],  # código repetido en el archivo
        ],
    )
    result = shop.send("products", content, confirm=True)
    assert result["committed"] is False
    assert [e["row"] for e in result["errors"]] == [3]
    fixed = xlsx(
        PRODUCT_HEADER, [["GLI-1", "Glifosato", "Insumo", "L", "Insumos", "", "0,05 bidón"]]
    )
    assert shop.send("products", fixed, confirm=True)["committed"] is True
    product = shop.c.get("/api/v1/products", params={"q": "Glifosato"}).json()["items"][0]
    assert product["code"] == "GLI-1"
    assert product["category"]["name"] == "Insumos"
    assert [c["unit"]["code"] for c in product["conversions"]] == ["bidón"]


def test_parties_import_validates_cuit(shop: Importer) -> None:
    header = [
        "Razón social *",
        "CUIT",
        "Condición IVA *",
        "Es cliente",
        "Es proveedor",
        "Días de pago",
    ]
    content = xlsx(
        header,
        [
            ["Cliente Uno", "20-12345678-6", "Monotributista", "Sí", "No", 15],
            ["Cliente Dos", "20-11111111-1", "Consumidor Final", "Sí", "", ""],  # CUIT inválido
        ],
    )
    result = shop.send("parties", content)
    assert [e["row"] for e in result["errors"]] == [3]
    ok = xlsx(header, [["Cliente Uno", "20-12345678-6", "Monotributista", "Sí", "No", 15]])
    assert shop.send("parties", ok, confirm=True)["committed"] is True


def test_opening_stock_one_document_per_warehouse(shop: Importer) -> None:
    header = ["Almacén *", "Producto *", "Cantidad *", "Unidad", "Costo unitario *", "Fecha"]
    content = xlsx(
        header,
        [
            ["Depósito", "Insecticida", "5", "", "1200", "01/09/2026"],
            ["Depósito", shop_code(shop, shop.poison), "5", "L", "1200", "01/09/2026"],
        ],
    )
    result = shop.send("opening_stock", content, confirm=True)
    assert result["committed"] is True, result
    assert shop.stock(shop.poison) == Decimal(20)  # 10 del escenario + 10
    docs = shop.c.get("/api/v1/stock-documents", params={"q": "Stock inicial"}).json()["items"]
    assert len([d for d in docs if d["reference"] == "Stock inicial"]) == 1


def shop_code(shop: Importer, product_id: str) -> str:
    return shop.c.get(f"/api/v1/products/{product_id}").json()["code"]  # type: ignore[no-any-return]


def test_opening_balances_are_not_sales(shop: Importer) -> None:
    header = ["Tipo *", "Cliente/proveedor *", "Fecha *", "Vencimiento", "Comprobante", "Importe *"]
    content = xlsx(
        header,
        [
            [
                "Cliente",
                "Verdulería Juan",
                "01/08/2026",
                "01/09/2026",
                "FC B 0001-00000077",
                "150.000,50",
            ],
            ["Proveedor", "Agroinsumos SA", "10/08/2026", "", "", "-2000"],
        ],
    )
    result = shop.send("opening_balances", content, confirm=True)
    assert result["committed"] is True, result
    sales = shop.c.get("/api/v1/current-accounts", params={"direction": "sale"}).json()
    customer = next(b for b in sales if b["party"]["id"] == shop.customer)
    assert Decimal(customer["balance"]) == Decimal("150000.50")
    assert Decimal(customer["overdue"]) == Decimal("150000.50")
    purchases = shop.c.get("/api/v1/current-accounts", params={"direction": "purchase"}).json()
    supplier = next(b for b in purchases if b["party"]["id"] == shop.supplier)
    assert Decimal(supplier["advances"]) == Decimal(2000)
    report = shop.c.get("/api/v1/reports/sales").json()
    assert report["rows"] == []  # el saldo inicial no es una venta
    docs = shop.c.get(
        "/api/v1/commercial-documents", params={"direction": "sale", "only_pending": True}
    ).json()["items"]
    assert docs[0]["invoice_label"] == "Saldo inicial FC B 0001-00000077"


def test_wrong_file_and_permissions(shop: Importer, client: TestClient, db: Session) -> None:
    bad = client.post(
        "/api/v1/imports/products/preview", files={"file": ("x.xlsx", b"no es excel", "text/plain")}
    )
    assert bad.status_code == 409
    missing = client.post(
        "/api/v1/imports/products/preview",
        files={"file": ("x.xlsx", xlsx(["Nombre"], [["x"]]), "application/octet-stream")},
    )
    assert missing.json()["error"]["code"] == "HEADER"
    login_as(client, db, READ_ONLY, username="lector")
    assert client.get("/api/v1/imports").json() == []
    assert client.get("/api/v1/imports/products/template").status_code == 403
