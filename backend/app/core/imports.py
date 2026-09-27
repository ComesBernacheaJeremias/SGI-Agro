"""Importación desde Excel (ADR-021): una definición por tipo de dato → plantilla, vista previa
y confirmación.

- La vista previa y la confirmación corren **la misma importación** con los servicios reales
  (mismas validaciones que la pantalla), cada fila en su propio savepoint: un error en una fila
  no frena las demás y queda informado con su número de fila.
- Vista previa: al final se deshace todo. Confirmación: si hubo algún error no se guarda nada
  (todo o nada).
"""

import re
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.core.errors import AppError, BusinessRuleError
from app.core.permissions import Permission
from app.core.schemas import Schema

MAX_ROWS = 5000


class RowError(Exception):
    """Error de una fila (mensaje para el usuario)."""


def normalize(text: str) -> str:
    """Para comparar textos: sin tildes, minúsculas, espacios simples."""
    plain = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", plain).strip().lower()


# --- Conversión de celdas (formato argentino) ---


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def cell_decimal(value: Any, label: str) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    text = str(value).strip().replace("$", "").replace(" ", "")
    if "," in text:  # 1.234,56 → 1234.56
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise RowError(f"{label}: '{value}' no es un número.") from exc


def cell_date(value: Any, label: str) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value).strip(), "%d/%m/%Y").date()
    except ValueError as exc:
        raise RowError(f"{label}: '{value}' no es una fecha dd/mm/aaaa.") from exc


def cell_bool(value: Any) -> bool:
    return normalize(cell_text(value)) in {"si", "s", "x", "1", "true", "verdadero"}


def choice(value: Any, label: str, options: dict[Any, str]) -> Any:
    """Busca por etiqueta ("Reventa") o por valor ("resale")."""
    text = normalize(cell_text(value))
    for key, option_label in options.items():
        if text in (normalize(option_label), normalize(str(key))):
            return key
    valid = ", ".join(options.values())
    raise RowError(f"{label}: '{value}' no es válido (opciones: {valid}).")


def required(value: Any, label: str) -> Any:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise RowError(f"{label}: obligatorio.")
    return value


# --- Definición ---


@dataclass(frozen=True)
class ImportColumn:
    key: str
    label: str
    required: bool = False
    help: str = ""
    example: str = ""


RowHandler = Callable[[Session, dict[str, Any], dict[str, Any]], None]
Finisher = Callable[[Session, dict[str, Any]], None]


@dataclass(frozen=True)
class ImportDef:
    key: str
    title: str
    description: str
    permission: Permission
    columns: list[ImportColumn]
    handle_row: RowHandler  # valida/crea una fila; `context` acumula entre filas
    finish: Finisher | None = None  # crea lo agrupado (ej. un ingreso de stock por almacén)
    notes: list[str] = field(default_factory=list)


IMPORTS: dict[str, ImportDef] = {}


def define_import(definition: ImportDef) -> ImportDef:
    if definition.key in IMPORTS:
        raise ValueError(f"Importación duplicada: {definition.key}")
    IMPORTS[definition.key] = definition
    return definition


# --- Resultado ---


class ImportRowError(Schema):
    row: int  # número de fila en el Excel
    message: str


class ImportResult(Schema):
    total: int
    valid: int
    errors: list[ImportRowError]
    committed: bool


def _message(exc: Exception, labels: dict[str, str]) -> str:
    if isinstance(exc, ValidationError):
        parts = []
        for err in exc.errors():
            field_name = str(err["loc"][0]) if err["loc"] else ""
            parts.append(f"{labels.get(field_name, field_name)}: {err['msg']}")
        return "; ".join(parts)
    return str(exc.args[0]) if exc.args else str(exc)


def read_rows(definition: ImportDef, content: bytes) -> list[tuple[int, dict[str, Any]]]:
    """Filas de la primera hoja; la fila 1 son los encabezados de la plantilla."""
    try:
        sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except Exception as exc:  # archivo dañado o no es Excel
        raise BusinessRuleError("El archivo no es un Excel válido (.xlsx).", code="FILE") from exc
    rows = sheet.iter_rows(values_only=True)
    header = [normalize(cell_text(h)).rstrip(" *") for h in next(rows, ())]
    positions: dict[str, int] = {}
    for column in definition.columns:
        label = normalize(column.label)
        if label in header:
            positions[column.key] = header.index(label)
        elif column.required:
            raise BusinessRuleError(
                f"Falta la columna '{column.label}'. Usá la plantilla.", code="HEADER"
            )
    result = []
    for number, values in enumerate(rows, start=2):
        row = {key: values[i] if i < len(values) else None for key, i in positions.items()}
        if all(cell_text(v) == "" for v in row.values()):
            continue  # filas vacías
        result.append((number, row))
    if len(result) > MAX_ROWS:
        raise BusinessRuleError(f"Máximo {MAX_ROWS} filas por archivo.", code="TOO_MANY_ROWS")
    return result


def run_import(
    session: Session, definition: ImportDef, content: bytes, *, commit: bool
) -> ImportResult:
    rows = read_rows(definition, content)
    labels = {c.key: c.label for c in definition.columns}
    errors: list[ImportRowError] = []
    context: dict[str, Any] = {}
    outer = session.begin_nested()
    for number, values in rows:
        savepoint = session.begin_nested()
        try:
            definition.handle_row(session, {**values, "_row": number}, context)
            savepoint.commit()
        except (RowError, AppError, ValidationError) as exc:
            savepoint.rollback()
            errors.append(ImportRowError(row=number, message=_message(exc, labels)))
    if not errors and definition.finish:
        try:
            definition.finish(session, context)
        except RowError as exc:
            errors.append(ImportRowError(row=context.get("_failed_row", 0), message=str(exc)))
        except AppError as exc:
            errors.append(
                ImportRowError(row=context.get("_failed_row", 0), message=_message(exc, labels))
            )
    committed = commit and not errors
    if committed:
        outer.commit()
    else:
        outer.rollback()
    failed_rows = {e.row for e in errors}
    return ImportResult(
        total=len(rows),
        valid=len(rows) - len(failed_rows),
        errors=sorted(errors, key=lambda e: e.row),
        committed=committed,
    )


def template(definition: ImportDef) -> bytes:
    """Plantilla: hoja "Datos" con encabezados (* = obligatorio) e "Instrucciones"."""
    wb = Workbook()
    data = wb.active
    assert data is not None  # noqa: S101
    data.title = "Datos"
    data.append([f"{c.label} *" if c.required else c.label for c in definition.columns])
    for i, column in enumerate(definition.columns, start=1):
        cell = data.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6B3F")
        data.column_dimensions[get_column_letter(i)].width = max(len(column.label) + 4, 16)
    data.freeze_panes = "A2"

    help_sheet = wb.create_sheet("Instrucciones")
    help_sheet.append([definition.title])
    help_sheet["A1"].font = Font(bold=True, size=14)
    help_sheet.append([definition.description])
    help_sheet.append([])
    help_sheet.append(["Columna", "Obligatoria", "Qué va", "Ejemplo"])
    for cell in help_sheet[4]:
        cell.font = Font(bold=True)
    for column in definition.columns:
        help_sheet.append(
            [column.label, "Sí" if column.required else "", column.help, column.example]
        )
    for note in definition.notes:
        help_sheet.append([])
        help_sheet.append([note])
    for letter, width in zip("ABCD", (28, 12, 70, 28), strict=True):
        help_sheet.column_dimensions[letter].width = width
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
