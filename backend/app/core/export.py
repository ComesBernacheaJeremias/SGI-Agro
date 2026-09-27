"""Exportación de un `TableReport` a Excel (openpyxl) y PDF (reportlab), con formato argentino.

Excel guarda números y fechas reales (se pueden sumar/filtrar) con formato de celda;
el separador de miles/decimales lo pone Excel según la configuración regional.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.platypus.flowables import Flowable

from app.core.formatting import format_date, format_money, format_quantity
from app.core.reports import Cell, ColumnKind, ReportRow, TableReport

NUMERIC: set[ColumnKind] = {"money", "quantity", "percent"}
EXCEL_FORMATS: dict[ColumnKind, str] = {
    "money": "#,##0.00",
    "quantity": "#,##0.00#",
    "percent": "0.0%",
    "date": "dd/mm/yyyy",
    "text": "@",
}


def format_cell(value: Cell, kind: ColumnKind) -> str:
    """Texto de una celda (PDF); mismas reglas que la pantalla."""
    if value is None or value == "":
        return ""
    if isinstance(value, date):
        return format_date(value)
    if isinstance(value, Decimal):
        if kind == "money":
            return format_money(value)
        if kind == "percent":
            return f"{format_quantity(value)} %"
        return format_quantity(value)
    return str(value)


def _rows_with_totals(report: TableReport) -> list[ReportRow]:
    rows = list(report.rows)
    if report.totals:
        first = report.columns[0].key
        cells = {first: "Total", **{k: v for k, v in report.totals.items() if k != first}}
        rows.append(ReportRow(cells=cells, style="total"))
    return rows


# --- Excel ---


def to_xlsx(report: TableReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101
    ws.title = report.title[:31]
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([report.subtitle])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([c.label for c in report.columns])
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6B3F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [len(c.label) for c in report.columns]
    for row in _rows_with_totals(report):
        values: list[object] = []
        for i, col in enumerate(report.columns):
            value = row.cells.get(col.key)
            if isinstance(value, Decimal):
                values.append(float(value / 100 if col.kind == "percent" else value))
            elif isinstance(value, str) and i == 0 and row.indent:
                values.append("    " * row.indent + value)
            else:
                values.append(value)
            widths[i] = max(widths[i], len(format_cell(value, col.kind)))
        ws.append(values)
        for i, col in enumerate(report.columns, start=1):
            target = ws.cell(row=ws.max_row, column=i)
            target.number_format = EXCEL_FORMATS[col.kind]
            if row.style != "normal":
                target.font = Font(bold=True)

    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = f"A{header_row + 1}"
    for note in report.notes:
        ws.append([])
        ws.append([note])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- PDF ---


def _footer(canvas: object, doc: SimpleDocTemplate) -> None:
    canvas.saveState()  # type: ignore[attr-defined]
    canvas.setFont("Helvetica", 7)  # type: ignore[attr-defined]
    text = f"SGI Agro · generado el {format_date(date.today())} · página {doc.page}"
    canvas.drawRightString(doc.pagesize[0] - 12 * mm, 8 * mm, text)  # type: ignore[attr-defined]
    canvas.restoreState()  # type: ignore[attr-defined]


def to_pdf(report: TableReport) -> bytes:
    pagesize = landscape(A4) if len(report.columns) > 6 else A4
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=report.title,
    )
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    numeric = [c.kind in NUMERIC for c in report.columns]

    data: list[list[object]] = [
        [Paragraph(f"<b>{escape(c.label)}</b>", cell_style) for c in report.columns]
    ]
    table_style: list[Any] = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F6B3F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDDDDD")),
    ]
    for i, is_number in enumerate(numeric):
        if is_number:
            table_style.append(("ALIGN", (i, 1), (i, -1), "RIGHT"))
    for r, row in enumerate(_rows_with_totals(report), start=1):
        texts = [format_cell(row.cells.get(c.key), c.kind) for c in report.columns]
        if row.indent:
            texts[0] = "    " * row.indent + texts[0]
        bold = row.style != "normal"
        data.append(
            [
                t
                if numeric[i]
                else Paragraph(f"<b>{escape(t)}</b>" if bold else escape(t), cell_style)
                for i, t in enumerate(texts)
            ]
        )
        if bold:
            table_style.append(("FONTNAME", (0, r), (-1, r), "Helvetica-Bold"))
            table_style.append(("BACKGROUND", (0, r), (-1, r), colors.HexColor("#F1F5F1")))

    story: list[Flowable] = [
        Paragraph(escape(report.title), styles["Title"]),
        Paragraph(escape(report.subtitle), styles["Normal"]),
        Spacer(1, 4 * mm),
        Table(data, repeatRows=1, style=TableStyle(table_style)),
    ]
    if not report.rows:
        story.append(Paragraph("Sin datos para los filtros elegidos.", styles["Normal"]))
    for note in report.notes:
        story += [Spacer(1, 2 * mm), Paragraph(escape(note), styles["Italic"])]
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()
